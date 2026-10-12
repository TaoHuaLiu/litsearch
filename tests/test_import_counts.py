from openpyxl import load_workbook

from litsearch.enrich import dedupe
from litsearch.export import write_xlsx
from litsearch.importers import from_bibtex, from_openalex_md, from_ris, load
from litsearch.models import Record
from prisma_counts import counts

BIB = """@article{pop00001,
\tauthor = {A Smith and B Jones},
\ttitle = {Adaptive Transparency: A Study},
\tyear = {2025},
\tabstract = {… snippet …},
\tnote = {Query date: 2026-10-02}
}
"""

MD = """## Q

### 1. Adaptive transparency
EXCLUDE

- Auteurs : Ann Lee, Bob Ray
- Annee : 2024 — Type : article
- DOI : [https://doi.org/10.1/ABC](https://doi.org/10.1/ABC)
- Resume :

> A full abstract.
"""


def test_bibtex_snippet_flagged(tmp_path):
    p = tmp_path / "x.bibtex"; p.write_text(BIB, encoding="utf-8")
    (r,) = from_bibtex(p, "pop")
    assert r.title == "Adaptive Transparency: A Study" and r.year == 2025
    assert r.authors == ["A Smith", "B Jones"] and r.abstract_is_snippet and r.sources == {"pop"}


def test_openalex_md_and_exclude_marker_not_imported(tmp_path):
    p = tmp_path / "o.md"; p.write_text(MD, encoding="utf-8")
    (r,) = from_openalex_md(p, "oa")
    assert r.doi == "10.1/abc" and r.abstract == "A full abstract." and r.authors == ["Ann Lee", "Bob Ray"]


def test_ris(tmp_path):
    p = tmp_path / "e.ris"
    p.write_text("TY  - JOUR\nTI  - A title\nAU  - Doe, J\nAU  - Roe, K\nPY  - 2021///\nDO  - 10.2/x\nAB  - abs\nER  - \n")
    (r,) = from_ris(p, "ebsco")
    assert r.authors == ["Doe, J", "Roe, K"] and r.year == 2021 and r.doi == "10.2/x"


def test_dedupe_across_imports_doi_prefix_and_title_punctuation(tmp_path):
    a = Record(title="Adaptive Transparency: A Study", year=2025, doi="https://doi.org/10.1/ABC", sources={"pop"})
    b = Record(title="Adaptive transparency - a study", year=2025, sources={"scholar"})
    c = Record(title="Other", year=2020, doi="10.9/z", sources={"pop"})
    out = dedupe([a, b, c])
    assert len(out) == 2 and {"pop", "scholar"} == next(r for r in out if r.title.startswith("Adaptive")).sources


def test_counts_and_invariants(tmp_path):
    recs = [Record(title=f"T{i}", year=2024, sources={"s"}) for i in range(4)]
    stats = [{"label": "import", "query": "f", "source": "s", "count": 6}]
    p = tmp_path / "r.xlsx"; write_xlsx(p, recs, stats)
    wb = load_workbook(p); ws = wb["Results"]
    h = {c.value: c.column for c in ws[1]}
    ws.cell(2, h["Screening"], "exclure"); ws.cell(2, h["Motif_screening"], "optique")
    ws.cell(3, h["Screening"], "inclure"); ws.cell(3, h["Texte_integral"], "exclure")  # missing reason
    ws.cell(4, h["Screening"], "inclure"); ws.cell(4, h["Texte_integral"], "inclure")
    wb.save(p)
    c, problems = counts(str(p))
    assert (c["identified"], c["duplicates"], c["screened"]) == (6, 2, 4)
    assert (c["screening_excluded"], c["fulltext_assessed"], c["fulltext_excluded"], c["included"]) == (1, 2, 1, 1)
    assert c["screening_pending"] == 1 and len(problems) == 1 and "full text without reason" in problems[0]
