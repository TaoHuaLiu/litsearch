"""Simule le travail du relecteur : remplit les colonnes de décision du registre (ce que vous faites à la main
dans Excel). Usage : fill_decisions.py registre.xlsx registre_decide.xlsx"""
import sys

from openpyxl import load_workbook

# (début du titre) -> (Screening, Motif_screening, Texte_integral, Motif_FT)
DECISIONS = {
    "Smart windows": ("exclure", "transparence optique (hors sujet)", "", ""),
    "Government transparency": ("exclure", "transparence gouvernementale (hors sujet)", "", ""),
    "A review of human-autonomy": ("exclure", "publié avant 2015", "", ""),
    "Adaptive transparency for autonomous": ("inclure", "", "inclure", ""),
    "Trust calibration": ("inclure", "", "inclure", ""),
    "Explainable AI": ("inclure", "", "inclure", ""),
    "Workload-adaptive": ("douteux", "", "exclure", "pas d'évaluation de l'effet sur la confiance"),
}

wb = load_workbook(sys.argv[1])
ws = wb["Results"]
col = {c.value: c.column for c in ws[1]}
for row in range(2, ws.max_row + 1):
    title = ws.cell(row, col["Title"]).value or ""
    for start, vals in DECISIONS.items():
        if title.startswith(start):
            for name, v in zip(("Screening", "Motif_screening", "Texte_integral", "Motif_FT"), vals):
                ws.cell(row, col[name], v or None)
wb.save(sys.argv[2])
