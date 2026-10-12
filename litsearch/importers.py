"""Import results exported by other tools (Publish or Perish, Scholar xlsx, OpenAlex markdown, EBSCO RIS/CSV)
as Records, so they go through the same dedupe / abstract backfill / Match pipeline as live searches."""
from __future__ import annotations

import csv
import re
from pathlib import Path

from openpyxl import load_workbook

from .models import Record, clean_doi

# header aliases (lower-case) -> Record field, for xlsx / csv imports
ALIASES = {
    "title": "title", "titre": "title",
    "authors": "authors", "author": "authors", "auteurs": "authors",
    "year": "year", "année": "year", "annee": "year",
    "source": "publication", "publication": "publication",
    "revue / conférence / source": "publication",
    "doi": "doi",
    "articleurl": "url", "url": "url", "url publication": "url",
    "full abstract": "abstract", "abstract": "abstract", "résumé": "abstract", "resume": "abstract",
    "extrait google scholar (pas le résumé intégral)": "snippet",
}


def _year(v) -> int | None:
    m = re.search(r"\b(19|20)\d{2}\b", str(v or ""))
    return int(m.group()) if m else None


def _authors(v) -> list[str]:
    v = str(v or "")
    if ";" not in v and ", " in v:  # PoP style "F Wang, C Lopez": split unless it is a single "Last, F"
        parts = [p.strip() for p in v.split(",")]
        if all(" " in p for p in parts):
            return [p for p in parts if p]
    return [a.strip() for a in re.split(r";| and ", v) if a.strip()]


def _record(d: dict, source: str, snippet_only: bool = False) -> Record | None:
    title = re.sub(r"\s+", " ", str(d.get("title") or "")).strip()
    if not title:
        return None
    abstract = str(d.get("abstract") or "").strip()
    snippet = False
    if not abstract and d.get("snippet"):
        abstract, snippet = str(d["snippet"]).strip(), True
    elif snippet_only and abstract:
        snippet = True
    return Record(title=title, authors=_authors(d.get("authors")), year=_year(d.get("year")),
                  publication=str(d.get("publication") or "").strip(), abstract=abstract,
                  doi=clean_doi(str(d.get("doi") or "")), url=str(d.get("url") or "").strip(),
                  abstract_is_snippet=snippet, sources={source})


def _rows_to_records(header: list, rows, source: str) -> list[Record]:
    cols = {i: ALIASES[str(h).strip().lower()] for i, h in enumerate(header)
            if h is not None and str(h).strip().lower() in ALIASES}
    out = []
    for row in rows:
        d = {f: row[i] for i, f in cols.items() if i < len(row) and row[i] not in (None, "")}
        r = _record(d, source)
        if r:
            out.append(r)
    return out


def from_xlsx(path: Path, source: str) -> list[Record]:
    """Any sheet with a header row containing a Title/Titre column (header searched in the first 10 rows).
    Several sheets: only those with a title column are read; the same reference in several sheets is merged later."""
    wb = load_workbook(path, read_only=True, data_only=True)
    out: list[Record] = []
    for ws in wb:
        if ws.title.startswith("Thématique"):  # per-theme sheets duplicate "Références uniques"
            continue
        rows = list(ws.iter_rows(values_only=True))
        for i, row in enumerate(rows[:10]):
            if any(str(c or "").strip().lower() in ("title", "titre") for c in row):
                out += _rows_to_records(list(row), rows[i + 1:], source)
                break
    return out


def from_csv(path: Path, source: str) -> list[Record]:
    with open(path, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))
    return _rows_to_records(rows[0], rows[1:], source) if rows else []


def from_bibtex(path: Path, source: str) -> list[Record]:
    """Publish or Perish BibTeX. Its abstract field is a truncated Scholar snippet, flagged as such."""
    text = path.read_text(encoding="utf-8-sig")
    out = []
    for entry in re.split(r"(?m)^@\w+\{", text)[1:]:
        d = {}
        for m in re.finditer(r"(?m)^\s*(\w+)\s*=\s*\{(.*?)\},?\s*$", entry):
            d.setdefault(m.group(1).lower(), m.group(2))
        r = _record({"title": d.get("title"), "authors": (d.get("author") or "").replace(" and ", ";"),
                     "year": d.get("year"), "publication": d.get("journal") or d.get("booktitle"),
                     "doi": d.get("doi"), "url": d.get("url"), "abstract": d.get("abstract")},
                    source, snippet_only=True)
        if r:
            out.append(r)
    return out


def from_ris(path: Path, source: str) -> list[Record]:
    """EBSCO / generic RIS export."""
    out, cur = [], {}

    def flush():
        if cur:
            r = _record({"title": cur.get("TI") or cur.get("T1"), "authors": ";".join(cur.get("AU", [])),
                         "year": cur.get("PY") or cur.get("Y1"), "publication": cur.get("JO") or cur.get("T2") or cur.get("JF"),
                         "doi": cur.get("DO"), "url": cur.get("UR"), "abstract": cur.get("AB") or cur.get("N2")}, source)
            if r:
                out.append(r)
        cur.clear()

    for line in path.read_text(encoding="utf-8-sig", errors="replace").splitlines():
        m = re.match(r"^([A-Z][A-Z0-9])  - ?(.*)$", line)
        if not m:
            continue
        tag, val = m.groups()
        if tag == "ER":
            flush()
        elif tag == "AU":
            cur.setdefault("AU", []).append(val.strip())
        else:
            cur.setdefault(tag, val.strip())
    flush()
    return out


def from_openalex_md(path: Path, source: str) -> list[Record]:
    """`openalex_results.md`: one `### n. title` block per record. A bare `EXCLUDE` line under the title
    (manual annotation) is not imported: screening decisions belong in the decision columns."""
    out = []
    for block in re.split(r"(?m)^### \d+\. ", path.read_text(encoding="utf-8"))[1:]:
        title = block.splitlines()[0].strip()
        field = lambda k: (re.search(rf"(?m)^- {k} : (.*)$", block) or [None, ""])[1].strip()
        doi = re.search(r"https://doi\.org/[^\s)\]]+", field("DOI"))
        ab = re.search(r"(?m)^- Resume : *\n\n> (.*?)(?:\n\n|\Z)", block, re.S)
        r = _record({"title": title, "authors": field("Auteurs").replace(", ", ";"), "year": field("Annee"),
                     "doi": doi.group() if doi else "", "url": field("OpenAlex"),
                     "abstract": ab.group(1).replace("\n> ", " ") if ab else ""}, source)
        if r:
            out.append(r)
    return out


READERS = {".xlsx": from_xlsx, ".csv": from_csv, ".bib": from_bibtex, ".bibtex": from_bibtex,
           ".ris": from_ris, ".md": from_openalex_md}


def load(path: str | Path, source: str | None = None) -> list[Record]:
    p = Path(path)
    reader = READERS.get(p.suffix.lower())
    if not reader:
        raise ValueError(f"unsupported import format: {p.suffix} (use {', '.join(READERS)})")
    return reader(p, source or p.stem)
