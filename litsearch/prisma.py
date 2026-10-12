"""PRISMA 2020 counts from a litsearch xlsx (sheets "Results" and "Queries").

    litsearch-prisma results.xlsx [--csv prisma_counts.csv]   (or: python prisma_counts.py ...)

Identification = sum of "Raw results" in Queries (live searches and imports). Duplicates = identified - unique rows.
Screening column: inclure / exclure / douteux. Texte_integral column: inclure / exclure (for rows that passed screening).
Exits non-zero if an invariant is broken (missing reason, full-text decision on an excluded row, ...)."""
from __future__ import annotations

import argparse
import csv
import sys
from collections import Counter

from openpyxl import load_workbook


def counts(path: str) -> tuple[dict, list[str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    res = list(wb["Results"].iter_rows(values_only=True))
    head = res[0]
    col = {h: i for i, h in enumerate(head)}
    for need in ("Screening", "Motif_screening", "Texte_integral", "Motif_FT"):
        if need not in col:
            raise SystemExit(f"column {need!r} missing: re-export with the current litsearch")
    rows = [r for r in res[1:] if r[col["Title"]]]
    cell = lambda r, h: str(r[col[h]] or "").strip().lower()

    by_source: Counter = Counter()
    for r in wb["Queries"].iter_rows(min_row=2, values_only=True):
        if r[0] and isinstance(r[3], (int, float)):
            by_source[r[2]] += int(r[3])
    identified = sum(by_source.values())

    problems: list[str] = []
    screened = len(rows)
    s_excl = [r for r in rows if cell(r, "Screening") == "exclure"]
    s_pending = [r for r in rows if cell(r, "Screening") == ""]
    ft = [r for r in rows if cell(r, "Screening") in ("inclure", "douteux")]
    ft_excl = [r for r in ft if cell(r, "Texte_integral") == "exclure"]
    ft_incl = [r for r in ft if cell(r, "Texte_integral") == "inclure"]
    ft_pending = [r for r in ft if cell(r, "Texte_integral") == ""]

    for r in s_excl:
        if not cell(r, "Motif_screening"):
            problems.append(f"excluded at screening without reason: {r[col['Title']][:70]}")
    for r in ft_excl:
        if not cell(r, "Motif_FT"):
            problems.append(f"excluded at full text without reason: {r[col['Title']][:70]}")
    for r in rows:
        if cell(r, "Screening") == "exclure" and cell(r, "Texte_integral"):
            problems.append(f"full-text decision on a row excluded at screening: {r[col['Title']][:70]}")
    if len(s_excl) + len(ft) + len(s_pending) != screened:
        problems.append("screening decisions do not add up (unknown value in Screening column?)")
    if len(ft_excl) + len(ft_incl) + len(ft_pending) != len(ft):
        problems.append("full-text decisions do not add up (unknown value in Texte_integral column?)")

    reasons = Counter(cell(r, "Motif_FT") for r in ft_excl)
    return {"by_source": dict(by_source), "identified": identified, "duplicates": identified - screened,
            "screened": screened, "screening_excluded": len(s_excl), "screening_pending": len(s_pending),
            "fulltext_assessed": len(ft), "fulltext_excluded": len(ft_excl), "fulltext_pending": len(ft_pending),
            "included": len(ft_incl), "fulltext_reasons": dict(reasons)}, problems


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(prog="litsearch-prisma")
    ap.add_argument("xlsx")
    ap.add_argument("--csv", help="write the counts as a two-column CSV (item,value)")
    a = ap.parse_args(argv)
    c, problems = counts(a.xlsx)
    print("IDENTIFICATION")
    for s, n in c["by_source"].items():
        print(f"  {s}: {n}")
    print(f"  total identified: {c['identified']}\n  duplicates removed: {c['duplicates']}")
    print(f"SCREENING\n  screened: {c['screened']}\n  excluded: {c['screening_excluded']}  (pending: {c['screening_pending']})")
    print(f"ELIGIBILITY\n  full texts assessed: {c['fulltext_assessed']}\n  excluded: {c['fulltext_excluded']}  (pending: {c['fulltext_pending']})")
    for m, n in c["fulltext_reasons"].items():
        print(f"    - {m or '(no reason)'}: {n}")
    print(f"INCLUDED: {c['included']}")
    if a.csv:
        with open(a.csv, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["item", "value"])
            for k in ("identified", "duplicates", "screened", "screening_excluded", "fulltext_assessed",
                      "fulltext_excluded", "included"):
                w.writerow([k, c[k]])
            for m, n in c["fulltext_reasons"].items():
                w.writerow([f"fulltext_excluded: {m}", n])
    for p in problems:
        print("! " + p, file=sys.stderr)
    return 1 if problems else 0


if __name__ == "__main__":
    raise SystemExit(main())
