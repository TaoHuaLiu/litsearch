from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import Record

COLUMNS = ["Author", "Title", "Year", "Publication", "Full Abstract", "DOI", "DOI_source", "DOI_confiance", "URL", "Sources", "Query", "Match",
           "Screening", "Motif_screening", "Texte_integral", "Motif_FT", "Date"]
WIDTHS = [32, 55, 8, 32, 90, 28, 12, 14, 38, 22, 12, 9, 12, 28, 14, 28, 12]
MAX_CELL = 32000


def _row(r: Record) -> list:
    ab = r.abstract
    if r.abstract_is_snippet and ab:
        ab = "[Scholar snippet only] " + ab
    return ["; ".join(a for a in r.authors if a), r.title, r.year or "", r.publication, ab[:MAX_CELL],
            r.doi, r.doi_source, r.doi_confidence, r.url, ", ".join(sorted(r.sources)), ", ".join(r.queries), r.match, "", "", "", "", ""]


def write_xlsx(path: Path, records: list[Record], query_stats: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(COLUMNS)
    for r in records:
        ws.append(_row(r))
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=c.column in (1, 2, 5), vertical="top")
    from openpyxl.worksheet.datavalidation import DataValidation
    for name, values in (("Screening", "inclure,exclure,douteux"), ("Texte_integral", "inclure,exclure")):
        col = ws.cell(1, COLUMNS.index(name) + 1).column_letter
        dv = DataValidation(type="list", formula1=f'"{values}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{max(ws.max_row, 2)}")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    qs = wb.create_sheet("Queries")
    qs.append(["Label", "Query", "Source", "Raw results", "Note"])
    for c in qs[1]:
        c.font = Font(bold=True)
    for s in query_stats:
        qs.append([s["label"], s["query"], s["source"], s["count"], s.get("note", "")])
    qs.column_dimensions["B"].width = 100
    qs.column_dimensions["E"].width = 60
    wb.save(path)


def write_csv(path: Path, records: list[Record]) -> None:
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(COLUMNS)
        for r in records:
            w.writerow(_row(r))
